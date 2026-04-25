import os
import io
import re
import uuid
import json
import base64
import tempfile
from http import HTTPStatus
from contextlib import contextmanager
from urllib.parse import quote

import bcrypt
import sqlite3
import dashscope
from dashscope import Generation, MultiModalConversation
from dashscope.audio.tts import SpeechSynthesizer
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, FileResponse
from pydantic import BaseModel
from dotenv import load_dotenv

from ai_tutor import AITutor
from worksheet_utils import create_worksheet_docx

load_dotenv()
dashscope.api_key = os.getenv("DASHSCOPE_API_KEY")

app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

DB_PATH = os.getenv("DB_PATH", "ai_tutor_users.db")

@contextmanager
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()

# 初始化数据库表
with get_db() as conn:
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users (
        username TEXT PRIMARY KEY,
        password_hash TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS chat_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT,
        role TEXT,
        content TEXT,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (username) REFERENCES users(username)
    )''')
    # 新增错题表
    c.execute('''CREATE TABLE IF NOT EXISTS mistakes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        session_id TEXT,
        subject TEXT,
        grade_level TEXT,
        question TEXT,
        wrong_answer TEXT,
        correct_answer TEXT,
        knowledge_point TEXT,
        error_type TEXT,
        reviewed INTEGER DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )''')

class ChatRequest(BaseModel):
    message: str
    username: str = "游客"
    logged_in: bool = False

class LoginRequest(BaseModel):
    username: str
    password: str

class RegisterRequest(BaseModel):
    username: str
    password: str

class HomeworkCheckRequest(BaseModel):
    content: str

class SimilarRequest(BaseModel):
    question: str
    correct_answer: str

class ChatMultimodalRequest(BaseModel):
    message: str
    image_base64: str = None

class VisualTeacherPlainRequest(BaseModel):
    image_base64: str
    question_text: str = ""

def save_chat_message(username, role, content):
    with get_db() as conn:
        c = conn.cursor()
        c.execute("INSERT INTO chat_history (username, role, content) VALUES (?, ?, ?)",
                  (username, role, content))
        conn.commit()

def load_chat_history(username, limit=50):
    with get_db() as conn:
        c = conn.cursor()
        c.execute("SELECT role, content FROM chat_history WHERE username = ? ORDER BY timestamp ASC LIMIT ?",
                  (username, limit))
        rows = c.fetchall()
        return [{"role": r['role'], "content": r['content']} for r in rows]

def clear_chat_history(username):
    with get_db() as conn:
        c = conn.cursor()
        c.execute("DELETE FROM chat_history WHERE username = ?", (username,))
        conn.commit()

# 数据库保存错题
def save_mistake(session_id, subject, grade_level, question, wrong_answer, correct_answer, knowledge_point, error_type):
    with get_db() as conn:
        c = conn.cursor()
        c.execute(
            "INSERT INTO mistakes (session_id, subject, grade_level, question, wrong_answer, correct_answer, knowledge_point, error_type) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (session_id, subject, grade_level, question, wrong_answer, correct_answer, knowledge_point, error_type)
        )
        conn.commit()

# 从数据库获取所有错题
def get_all_mistakes():
    with get_db() as conn:
        c = conn.cursor()
        c.execute("SELECT * FROM mistakes ORDER BY created_at DESC")
        return c.fetchall()

tutor = AITutor()

def call_qwen_vl(image_bytes, prompt="请提取图片中的所有题目文字，只返回文字内容。"):
    encoded = base64.b64encode(image_bytes).decode('utf-8')
    messages = [
        {
            "role": "user",
            "content": [
                {"image": f"data:image/png;base64,{encoded}"},
                {"text": prompt}
            ]
        }
    ]
    try:
        response = MultiModalConversation.call(
            model='qwen3.6-plus',
            messages=messages
        )
        if response.status_code == HTTPStatus.OK:
            return response.output.choices[0].message.content[0]["text"]
        else:
            return f"❌ 图片识别失败：{response.code} - {response.message}"
    except Exception as e:
        return f"❌ 视觉模型调用失败：{str(e)}"

def generate_speech_ali(text: str) -> bytes:
    result = SpeechSynthesizer.call(
        model='sambert-zhixiao-v1',
        text=text,
        sample_rate=16000,
        format='mp3'
    )
    if result.get_audio_data():
        return result.get_audio_data()
    else:
        raise Exception(f"TTS 失败：{result.get_response()}")

# ----- API 路由 -----
@app.post("/api/login")
def login(req: LoginRequest):
    with get_db() as conn:
        c = conn.cursor()
        c.execute("SELECT password_hash FROM users WHERE username = ?", (req.username,))
        row = c.fetchone()
        if row and bcrypt.checkpw(req.password.encode(), row['password_hash'].encode()):
            return {"success": True, "message": "登录成功"}
        return {"success": False, "message": "用户名或密码错误"}

@app.post("/api/register")
def register(req: RegisterRequest):
    hashed = bcrypt.hashpw(req.password.encode(), bcrypt.gensalt()).decode()
    with get_db() as conn:
        c = conn.cursor()
        try:
            c.execute("INSERT INTO users (username, password_hash) VALUES (?, ?)", (req.username, hashed))
            conn.commit()
            return {"success": True, "message": "注册成功"}
        except sqlite3.IntegrityError:
            return {"success": False, "message": "用户名已存在"}

@app.get("/api/history")
def get_history(username: str):
    return {"history": load_chat_history(username)}

@app.post("/api/clear-history")
def clear_history(username: str):
    clear_chat_history(username)
    return {"success": True}

@app.post("/api/chat")
def chat(req: ChatRequest):
    prompt = req.message.strip()
    if "作业" in prompt:
        pattern = r'(?:生成|出)(\d+|两|二|一|三|四|五|六|七|八|九|十)?道?[的]?(.+?)(?:作业|题目|问题)'
        match = re.search(pattern, prompt)
        if match:
            num = 5
            num_str = match.group(1)
            if num_str:
                chinese_map = {"一":1, "二":2, "两":2, "三":3, "四":4, "五":5, "六":6, "七":7, "八":8, "九":9, "十":10}
                if num_str.isdigit():
                    num = int(num_str)
                else:
                    num = chinese_map.get(num_str, 5)
            topic = match.group(2).strip()
            if topic:
                try:
                    ws = tutor.generate_worksheet(topic, "中等", num)
                    if ws and ws.get("questions"):
                        tutor.current_worksheet = ws
                        title = ws.get("title", topic)
                        questions_text = "\n".join([f"{i+1}. {q}" for i, q in enumerate(ws['questions'])])
                        reply = f"✅ **{title}** 已生成（共 {len(ws['questions'])} 题）：\n\n{questions_text}"
                        return {"reply": reply, "worksheet_ready": True, "title": title}
                    else:
                        reply = "❌ 生成失败，请稍后重试。"
                except Exception:
                    reply = "❌ 生成作业时发生错误。"
            else:
                reply = "请告诉我具体的知识点，例如：生成2道鸡兔同笼作业。"
        else:
            reply = "请使用类似“生成2道鸡兔同笼作业”的格式。"
    else:
        try:
            reply = tutor.chat(prompt)
        except Exception:
            reply = "❌ 老师暂时不在，请稍后再试。"
    return {"reply": reply, "worksheet_ready": False}

@app.get("/api/download-worksheet")
def download_worksheet(answers: bool = False):
    ws = getattr(tutor, "current_worksheet", None)
    if not ws:
        raise HTTPException(status_code=404, detail="没有可下载的作业，请先生成题目")
    try:
        title = ws.get("title", "练习题")
        questions = ws.get("questions", [])
        ans = ws.get("answers", [])
        docx = create_worksheet_docx(title, questions, ans, include_answers=answers)
        encoded_title = quote(title)
        return Response(
            content=docx,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=utf-8''{encoded_title}.docx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成文档失败: {str(e)}")

@app.post("/api/upload-homework")
async def upload_homework(file: UploadFile = File(...)):
    try:
        content = await file.read()
        text = ""
        if file.filename.endswith(('.png', '.jpg', '.jpeg', '.bmp')):
            text = call_qwen_vl(content)
        elif file.filename.endswith('.txt'):
            text = content.decode('utf-8')
        elif file.filename.endswith('.docx'):
            from docx import Document
            doc = Document(io.BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs])
        else:
            return {"reply": "❌ 不支持的文件格式，请上传图片、txt 或 docx"}

        if not text.strip() or text.startswith("❌"):
            return {"reply": f"❌ 内容提取失败：{text}"}

        result = tutor.grade_homework(subject="通用", grade_level="通用", homework_content=text)
        if result:
            score = result.get("score", 0)
            comment = result.get("comment", "加油！")
            msg = f"📋 **得分**：{score} / {result.get('total_score', 100)}\n**评语**：{comment}\n\n"

            mistakes_list = result.get("mistakes", [])
            if mistakes_list:
                msg += "**❌ 错题详情**：\n"
                for i, m in enumerate(mistakes_list, 1):
                    q = m.get("question", "")
                    correct = m.get("correct_answer", "")
                    sol = m.get("solution", "")
                    msg += f"{i}. {q} → 正确：{correct}\n"
                    if sol:
                        msg += f"   **解题步骤**：{sol}\n"
                    # 持久化错题到数据库
                    save_mistake(
                        session_id=tutor.session_id,
                        subject="通用",
                        grade_level="通用",
                        question=q,
                        wrong_answer=m.get("wrong_answer", ""),
                        correct_answer=correct,
                        knowledge_point=m.get("knowledge_point", ""),
                        error_type=m.get("error_type", "未分类")
                    )
            else:
                msg += "✅ 所有题目都正确，太棒了！\n"
                all_sols = result.get("all_solutions", [])
                if all_sols:
                    msg += "\n📝 **解题过程复习**：\n"
                    for sol in all_sols:
                        msg += f"**题目**：{sol.get('question', '')}\n**解法**：{sol.get('solution', '')}\n\n"
            return {"reply": msg}
        else:
            return {"reply": "✅ 批改完成，未发现明显错误。"}
    except Exception as e:
        return {"reply": f"❌ 批改失败: {str(e)}"}

@app.post("/api/photo-search")
async def photo_search(file: UploadFile = File(...)):
    try:
        content = await file.read()
        if not file.filename.endswith(('.png', '.jpg', '.jpeg', '.bmp')):
            return {"reply": "📷 仅支持图片题目哦"}
        text = call_qwen_vl(content, "请提取图片中的题目文字。")
        if not text.strip() or text.startswith("❌"):
            return {"reply": f"🔍 未能识别到题目：{text}"}
        reply = tutor.chat(f"请解答以下题目：\n{text}")
        return {"reply": reply}
    except Exception as e:
        return {"reply": f"📷 搜题失败: {str(e)}"}

# 错题本 API（从数据库读取）
@app.get("/api/mistakes")
def get_mistakes():
    try:
        mistakes = get_all_mistakes()
        result = []
        for m in mistakes:
            result.append({
                "id": m["id"],
                "question": m["question"],
                "correct_answer": m["correct_answer"],
                "error_type": m["error_type"],
                "reviewed": m["reviewed"],
                "created_at": m["created_at"]
            })
        return {"mistakes": result}
    except Exception:
        return {"mistakes": []}

@app.post("/api/mistakes/{mistake_id}/review")
def mark_reviewed(mistake_id: int):
    try:
        with get_db() as conn:
            c = conn.cursor()
            c.execute("UPDATE mistakes SET reviewed = 1 WHERE id = ?", (mistake_id,))
            conn.commit()
        return {"success": True}
    except Exception:
        return {"success": False}

@app.post("/api/mistakes/similar")
def generate_similar(req: SimilarRequest):
    try:
        new_q = tutor.generate_similar_question(req.question, req.correct_answer)
        return {"question": new_q}
    except Exception:
        return {"question": "生成失败，请重试"}

@app.get("/api/report")
def get_report():
    try:
        mistakes = get_all_mistakes()
        if not mistakes:
            return {"report": "🎉 本周没有错题，表现优秀！继续保持哦～"}
        mistakes_text = "\n".join([f"- {m['question']}" for m in mistakes])
        prompt = f"根据以下错题，生成一份简短的周学习报告：\n{mistakes_text}"
        report = tutor.chat(prompt)
        return {"report": report}
    except Exception:
        return {"report": "生成报告时出错"}

@app.post("/api/chat-multimodal")
def chat_multimodal(req: ChatMultimodalRequest):
    prompt = req.message.strip()
    image_b64 = req.image_base64
    if image_b64:
        try:
            response = MultiModalConversation.call(
                model='qwen3.6-plus',
                messages=[{
                    "role": "user",
                    "content": [
                        {"image": f"data:image/png;base64,{image_b64}"},
                        {"text": prompt if prompt else "请描述这张图片"}
                    ]
                }]
            )
            if response.status_code == HTTPStatus.OK:
                reply = response.output.choices[0].message.content[0]["text"]
            else:
                reply = f"❌ 多模态调用失败：{response.code}"
        except Exception as e:
            reply = f"❌ 多模态调用异常：{str(e)}"
    else:
        try:
            reply = tutor.chat(prompt)
        except Exception:
            reply = "❌ 老师暂时不在，请稍后再试。"
    return {"reply": reply}

# 超拟人老师讲解
@app.post("/api/visual-teacher-plain")
async def visual_teacher_plain(req: VisualTeacherPlainRequest):
    image_b64 = req.image_base64
    image_bytes = base64.b64decode(image_b64)
    ocr_text = call_qwen_vl(image_bytes, "请提取图片中的所有题目文字，只输出题目本身，不要解释。")
    if ocr_text.startswith("❌"):
        return {"error": f"题目提取失败：{ocr_text}"}

    full_prompt = (
        f"你是一位耐心的AI家教。这是一道题目：\n{ocr_text}\n\n"
        "请一步一步详细讲解这道题目的解题思路和计算过程，用中文，语气亲切活泼。\n"
        "每步讲解请包含完整的句子，至少2-3句话，详细解释这一步在做什么以及为什么这样做。\n"
        "请严格按照以下JSON格式输出（不要输出其他内容）：\n"
        "{\n"
        '  "steps": [\n'
        '    {"text": "第一步：首先我们分析题目给出的已知条件……"},\n'
        '    {"text": "第二步：接下来我们需要应用公式……因为……所以……"},\n'
        '    {"text": "第三步：最后我们将数值代入计算，得出最终答案……"}\n'
        '  ]\n'
        "}\n"
        "确保每步的text字段包含完整的讲解段落，不仅仅是单个短语或数字。"
    )
    try:
        response_text = tutor.chat(full_prompt)
        clean = response_text.replace("```json", "").replace("```", "").strip()
        steps_data = json.loads(clean)
        steps = steps_data.get("steps", [])
    except Exception as e:
        return {"error": f"讲解生成失败：{str(e)}"}

    return {"steps": steps}

# 逐句 TTS 端点
@app.post("/api/tts")
async def tts(text: str = Form(...)):
    try:
        audio_bytes = generate_speech_ali(text)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".mp3")
        tmp.write(audio_bytes)
        tmp.close()
        return {"url": f"http://127.0.0.1:8000/api/tts-audio?file={os.path.basename(tmp.name)}"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/tts-audio")
async def tts_audio(file: str):
    path = os.path.join(tempfile.gettempdir(), file)
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="音频文件不存在")
    return FileResponse(path, media_type="audio/mpeg")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)